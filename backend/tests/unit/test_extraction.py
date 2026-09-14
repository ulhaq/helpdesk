import io

import pytest
from docx import Document

from src.helpdesk.enums import HelpdeskErrorCode, TextFormat
from src.helpdesk.extraction import extract_text
from src.platform.core.exceptions import ClientException


def _pdf(text: str) -> bytes:
    """A minimal one-page PDF with a text layer."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, xref)
    )
    return out.getvalue()


def _docx() -> bytes:
    document = Document()
    document.add_heading("Refund policy", level=1)
    document.add_paragraph("Refunds within 30 days.")
    document.add_paragraph("")
    document.add_heading("Exceptions", level=2)
    document.add_paragraph("Damaged goods: always refund.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "VIP"
    table.rows[0].cells[1].text = "60 days"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_word_headings_become_markdown_headings() -> None:
    assert extract_text("policy.docx", _docx()) == (
        "# Refund policy\n\nRefunds within 30 days.\n\n## Exceptions\n\n"
        "Damaged goods: always refund.\n\nVIP | 60 days",
        TextFormat.MARKDOWN,
    )


def test_pdf_text_layer() -> None:
    text, text_format = extract_text("Terms.PDF", _pdf("Refunds within 30 days"))
    assert text == "Refunds within 30 days"
    assert text_format == TextFormat.PLAIN


def test_text_files() -> None:
    assert extract_text("notes.txt", "Café crème".encode("cp1252")) == (
        "Café crème",
        TextFormat.PLAIN,
    )
    assert extract_text("notes.txt", b"a  \r\n\r\n\r\n\r\nb\x00") == (
        "a\n\nb",
        TextFormat.PLAIN,
    )
    assert extract_text("guide.md", b"Intro only")[1] == TextFormat.MARKDOWN
    assert extract_text("notes.txt", b"## Refunds\n\n30 days")[1] == (
        TextFormat.MARKDOWN
    )


@pytest.mark.parametrize(
    ("filename", "data", "status_code", "error_code"),
    [
        ("sheet.xlsx", b"data", 415, HelpdeskErrorCode.UNSUPPORTED_FILE_TYPE),
        ("broken.pdf", b"not a pdf", 422, HelpdeskErrorCode.FILE_UNREADABLE),
        ("broken.docx", b"not a zip", 422, HelpdeskErrorCode.FILE_UNREADABLE),
    ],
)
def test_rejected_files(
    filename: str, data: bytes, status_code: int, error_code: HelpdeskErrorCode
) -> None:
    with pytest.raises(ClientException) as raised:
        extract_text(filename, data)
    assert raised.value.status_code == status_code
    assert raised.value.error_code == error_code
