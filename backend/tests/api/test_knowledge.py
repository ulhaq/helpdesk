import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.helpdesk.config import settings

_URL = "/v1/knowledge/documents"
_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _create(
    client: TestClient,
    title: str = "Refund exceptions",
    text: str = "VIP customers may get refunds after 60 days.",
) -> dict:
    response = client.post(_URL, json={"title": title, "text": text})
    assert response.status_code == 201, response.text
    return response.json()


def _upload(
    client: TestClient,
    filename: str,
    data: bytes,
    content_type: str = "text/plain",
    **form: str,
):
    return client.post(
        f"{_URL}/upload", files={"file": (filename, data, content_type)}, data=form
    )


def test_create_and_read_a_pasted_document(admin_authenticated: TestClient) -> None:
    document = _create(admin_authenticated)
    assert document["source"] == "text"
    assert document["format"] == "plain"
    assert document["filename"] is None
    assert document["text"] == "VIP customers may get refunds after 60 days."

    listed = admin_authenticated.get(_URL).json()
    assert [item["title"] for item in listed["items"]] == ["Refund exceptions"]
    assert "text" not in listed["items"][0]

    response = admin_authenticated.get(f"{_URL}/{document['id']}")
    assert response.json()["text"] == document["text"]


def test_pasted_markdown_is_split_at_headings(admin_authenticated: TestClient) -> None:
    document = _create(admin_authenticated, text="## Refunds\n\nWithin 30 days.")
    assert document["format"] == "markdown"


def test_upload_a_word_document(admin_authenticated: TestClient) -> None:
    word = Document()
    word.add_heading("Escalations", level=1)
    word.add_paragraph("Call the on-call lead.")
    buffer = io.BytesIO()
    word.save(buffer)
    data = buffer.getvalue()

    response = _upload(admin_authenticated, "escalations.docx", data, _DOCX)

    assert response.status_code == 201, response.text
    document = response.json()
    assert document["title"] == "escalations"
    assert document["source"] == "file"
    assert document["filename"] == "escalations.docx"
    assert document["content_type"] == _DOCX
    assert document["size_bytes"] == len(data)
    assert document["format"] == "markdown"
    assert document["text"] == "# Escalations\n\nCall the on-call lead."


def test_upload_with_a_title(admin_authenticated: TestClient) -> None:
    response = _upload(
        admin_authenticated, "script.txt", b"Greet, verify, help.", title="Phone script"
    )
    assert response.status_code == 201, response.text
    assert response.json()["title"] == "Phone script"


def test_rejected_uploads(
    admin_authenticated: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    for filename, data, status_code, error_code in [
        ("sheet.xlsx", b"data", 415, "unsupported_file_type"),
        ("empty.txt", b"  \n ", 422, "document_empty"),
        ("broken.pdf", b"not a pdf", 422, "file_unreadable"),
    ]:
        response = _upload(admin_authenticated, filename, data)
        assert response.status_code == status_code, filename
        assert response.json()["error_code"] == error_code

    monkeypatch.setattr(settings, "knowledge_max_upload_bytes", 10)
    response = _upload(admin_authenticated, "big.txt", b"x" * 11)
    assert response.status_code == 413
    assert response.json()["error_code"] == "file_too_large"
    assert admin_authenticated.get(_URL).json()["total"] == 0


def test_edit_and_delete_a_document(admin_authenticated: TestClient) -> None:
    document = _create(admin_authenticated)
    url = f"{_URL}/{document['id']}"

    response = admin_authenticated.patch(
        url, json={"title": "VIP refunds", "text": "## VIP\n\nUp to 90 days."}
    )
    assert response.status_code == 200, response.text
    updated = response.json()
    assert (updated["title"], updated["format"]) == ("VIP refunds", "markdown")

    response = admin_authenticated.patch(url, json={"text": "   "})
    assert response.json()["error_code"] == "document_empty"

    assert admin_authenticated.delete(url).status_code == 204
    assert admin_authenticated.get(url).status_code == 404


def test_documents_require_the_knowledge_base_permission(
    no_roles_authenticated: TestClient,
) -> None:
    assert no_roles_authenticated.get(_URL).status_code == 403
    response = no_roles_authenticated.post(_URL, json={"title": "x", "text": "y"})
    assert response.status_code == 403


def test_documents_are_scoped_to_the_organization(
    admin_authenticated: TestClient, organization2_admin_authenticated: TestClient
) -> None:
    document = _create(admin_authenticated)

    other = organization2_admin_authenticated
    assert other.get(f"{_URL}/{document['id']}").status_code == 404
    assert other.get(_URL).json()["total"] == 0
    assert other.delete(f"{_URL}/{document['id']}").status_code == 404
